"""Batch ZIP and clinician-score ingestion for the benchmark mode.

The signal package resolution is deliberately delegated to ``eval_package``.
This module only adds strict multi-patient discovery and the benchmark score
workbook; it never changes the existing EEG/EMG/IMU decoders.
"""
from __future__ import annotations

import json
import math
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable, Dict, Iterable, Mapping, Optional

from openpyxl import load_workbook

from eval_package import EvalPackage, read_eval_package, safe_extract_zip

from .runner import build_evaluation_input
from .schemas import (
    BenchmarkClinicalScores,
    BenchmarkEvaluationInput,
    BenchmarkPatientInfo,
    BenchmarkRunConfig,
)


REQUIRED_SCORE_COLUMNS = (
    "patient_id",
    "fma_wrist",
    "fma_hand",
    "hand_mas",
    "brunnstrom_hand",
)
_FORBIDDEN_SCORE_COLUMNS = {"bi", "fma_ue", "fma_upper", "fma_upper_limb"}


class BenchmarkBatchValidationError(ValueError):
    """Raised when a batch cannot be matched one-to-one to clinical scores."""


@dataclass(frozen=True)
class BenchmarkPatientPackage:
    package: EvalPackage
    patient: BenchmarkPatientInfo
    clinical_scores: BenchmarkClinicalScores


@dataclass(frozen=True)
class BenchmarkBatch:
    root: Path
    patients: tuple[BenchmarkPatientPackage, ...]
    score_file: Path


def _normalized_header(value: Any) -> str:
    return str(value or "").strip().lower().replace("-", "_").replace(" ", "_")


def _finite_number(value: Any, field: str, *, nonnegative: bool = False) -> float:
    if isinstance(value, bool):
        raise BenchmarkBatchValidationError(f"{field}必须是数值")
    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        raise BenchmarkBatchValidationError(f"{field}必须是数值") from exc
    if not math.isfinite(number) or (nonnegative and number < 0):
        raise BenchmarkBatchValidationError(f"{field}必须是有限非负数")
    return number


def _mas_text(value: Any, field: str) -> str:
    text = str(value or "").strip()
    if text in {"0", "1", "1+", "2", "3", "4"}:
        return text
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        numeric = float(value)
        if numeric in {0.0, 1.0, 2.0, 3.0, 4.0}:
            return str(int(numeric))
    raise BenchmarkBatchValidationError(
        f"{field}必须为0、1、1+、2、3或4"
    )


def _int_score(value: Any, field: str, lower: int, upper: int) -> int:
    number = _finite_number(value, field)
    if not number.is_integer() or not lower <= int(number) <= upper:
        raise BenchmarkBatchValidationError(f"{field}必须为{lower}–{upper}整数")
    return int(number)


def read_clinical_scores(path: Path) -> Dict[str, BenchmarkClinicalScores]:
    """Read the exact clinician-score columns and reject duplicates/BI."""
    path = Path(path)
    if not path.is_file():
        raise BenchmarkBatchValidationError(f"未找到临床评分文件：{path.name}")
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        try:
            raw_headers = next(rows)
        except StopIteration as exc:
            raise BenchmarkBatchValidationError("clinical_scores.xlsx为空") from exc
        headers = [_normalized_header(value) for value in raw_headers]
        if any(not header for header in headers):
            raise BenchmarkBatchValidationError("clinical_scores.xlsx存在空列名")
        duplicates = sorted({header for header in headers if headers.count(header) > 1})
        if duplicates:
            raise BenchmarkBatchValidationError("clinical_scores.xlsx列名重复：" + "、".join(duplicates))
        forbidden = sorted(set(headers) & _FORBIDDEN_SCORE_COLUMNS)
        if forbidden:
            raise BenchmarkBatchValidationError(
                "Benchmark禁止使用旧/额外量表字段：" + "、".join(forbidden)
            )
        missing = [name for name in REQUIRED_SCORE_COLUMNS if name not in headers]
        if missing:
            raise BenchmarkBatchValidationError("clinical_scores.xlsx缺少字段：" + "、".join(missing))
        index = {header: position for position, header in enumerate(headers)}
        output: Dict[str, BenchmarkClinicalScores] = {}
        for row_number, raw_row in enumerate(rows, start=2):
            if not any(value not in (None, "") for value in raw_row):
                continue
            patient_id = str(raw_row[index["patient_id"]] or "").strip()
            if not patient_id:
                raise BenchmarkBatchValidationError(f"第{row_number}行patient_id为空")
            if patient_id in output:
                raise BenchmarkBatchValidationError(f"patient_id重复：{patient_id}")
            fma_wrist = _finite_number(
                raw_row[index["fma_wrist"]],
                f"第{row_number}行fma_wrist",
                nonnegative=True,
            )
            fma_hand = _finite_number(raw_row[index["fma_hand"]], f"第{row_number}行fma_hand")
            if not 0 <= fma_hand <= 20:
                raise BenchmarkBatchValidationError(f"第{row_number}行fma_hand必须在0–20")
            output[patient_id] = BenchmarkClinicalScores(
                fma_wrist=fma_wrist,
                fma_hand=fma_hand,
                hand_mas=_mas_text(raw_row[index["hand_mas"]], f"第{row_number}行hand_mas"),
                brunnstrom_hand=_int_score(
                    raw_row[index["brunnstrom_hand"]],
                    f"第{row_number}行brunnstrom_hand",
                    1,
                    6,
                ),
            )
        if not output:
            raise BenchmarkBatchValidationError("clinical_scores.xlsx没有有效评分行")
        return output
    finally:
        workbook.close()


def _manifest(path: Path) -> Dict[str, Any]:
    try:
        value = json.loads(path.read_text(encoding="utf-8-sig"))
    except (OSError, json.JSONDecodeError) as exc:
        raise BenchmarkBatchValidationError(f"manifest.json解析失败：{path}") from exc
    if not isinstance(value, dict):
        raise BenchmarkBatchValidationError(f"manifest.json必须是对象：{path}")
    return value


def _patient_from_package(package: EvalPackage, root: Path) -> BenchmarkPatientInfo:
    values = dict(package.patient_prefill)
    manifest = _manifest(root / "manifest.json")
    for key in ("patient_id", "name", "sex", "age", "diagnosis", "disease_days", "paralysis_side"):
        if manifest.get(key) not in (None, ""):
            values[key] = manifest[key]
    required = ("patient_id", "age", "sex", "diagnosis", "disease_days", "paralysis_side")
    missing = [key for key in required if values.get(key) in (None, "")]
    if missing:
        raise BenchmarkBatchValidationError(
            f"患者{values.get('patient_id') or root.name}资料缺少：" + "、".join(missing)
        )
    if not values.get("name"):
        values["name"] = values["patient_id"]
    try:
        # EvalPackage intentionally exposes institution-specific compatibility
        # fields (hospital_patient_id, clinical_age, clinical_mapping).  The
        # benchmark contract only accepts the canonical patient fields below;
        # do not leak acquisition metadata into the strict evaluation model.
        canonical = {
            key: values.get(key)
            for key in (
                "patient_id",
                "name",
                "sex",
                "age",
                "diagnosis",
                "disease_days",
                "paralysis_side",
            )
        }
        return BenchmarkPatientInfo(**canonical)
    except Exception as exc:  # noqa: BLE001
        raise BenchmarkBatchValidationError(
            f"患者{values.get('patient_id') or root.name}基本信息非法：{exc}"
        ) from exc


def _manifest_roots(root: Path) -> list[Path]:
    direct = root / "manifest.json"
    if direct.is_file():
        return [root]
    roots = sorted({path.parent for path in root.rglob("manifest.json")})
    if not roots:
        raise BenchmarkBatchValidationError("ZIP中未找到任何患者manifest.json")
    return roots


def read_benchmark_batch(root: Path, institution: str) -> BenchmarkBatch:
    """Resolve one or many existing evaluation packages plus one workbook."""
    root = Path(root)
    score_files = sorted(root.rglob("clinical_scores.xlsx"))
    if len(score_files) != 1:
        raise BenchmarkBatchValidationError(
            f"ZIP中必须且只能有一个clinical_scores.xlsx，当前{len(score_files)}个"
        )
    scores = read_clinical_scores(score_files[0])
    patients: list[BenchmarkPatientPackage] = []
    seen_patient_ids: set[str] = set()
    for package_root in _manifest_roots(root):
        package = read_eval_package(package_root, institution=institution)
        patient = _patient_from_package(package, package_root)
        if patient.patient_id in seen_patient_ids:
            raise BenchmarkBatchValidationError(f"患者patient_id重复：{patient.patient_id}")
        seen_patient_ids.add(patient.patient_id)
        if patient.patient_id not in scores:
            raise BenchmarkBatchValidationError(
                f"患者{patient.patient_id}没有clinical_scores.xlsx评分"
            )
        patients.append(
            BenchmarkPatientPackage(
                package=package,
                patient=patient,
                clinical_scores=scores[patient.patient_id],
            )
        )
    extra = sorted(set(scores) - seen_patient_ids)
    if extra:
        raise BenchmarkBatchValidationError(
            "clinical_scores.xlsx存在找不到患者资料的patient_id：" + "、".join(extra)
        )
    return BenchmarkBatch(root=root, patients=tuple(patients), score_file=score_files[0])


def read_benchmark_batch_zip(
    zip_path: Path,
    dest_dir: Path,
    institution: str,
) -> BenchmarkBatch:
    extracted = safe_extract_zip(Path(zip_path), Path(dest_dir))
    return read_benchmark_batch(extracted, institution)


def prepare_batch_inputs(
    batch: BenchmarkBatch,
    *,
    biomarker_extractor: Callable[..., Mapping[str, Any]],
    config: Optional[BenchmarkRunConfig] = None,
) -> list[BenchmarkEvaluationInput]:
    """Compute biomarker inputs with the existing extractor for every patient."""
    inputs: list[BenchmarkEvaluationInput] = []
    for item in batch.patients:
        biomarkers = biomarker_extractor(
            item.package.eeg_paths,
            item.package.emg_paths,
            hand_function_stage=item.clinical_scores.brunnstrom_hand,
            affected_side=item.patient.paralysis_side,
            institution=item.package.institution,
        )
        inputs.append(
            build_evaluation_input(
                patient=item.patient,
                clinical_scores=item.clinical_scores,
                biomarkers=biomarkers,
                quality={"status": "pass", "source": "benchmark_batch"},
                config=config,
            )
        )
    return inputs


def existing_biomarker_extractor() -> Callable[..., Mapping[str, Any]]:
    """Return the same backend extractor used by the normal inference worker."""
    from inference import _load_backend_biomarkers

    return _load_backend_biomarkers().extract


def prepare_single_input(
    *,
    patient: BenchmarkPatientInfo,
    clinical_scores: BenchmarkClinicalScores,
    eeg_paths: Iterable[Path],
    emg_paths: Iterable[Path],
    institution: str,
    config: Optional[BenchmarkRunConfig] = None,
    biomarker_extractor: Optional[Callable[..., Mapping[str, Any]]] = None,
) -> BenchmarkEvaluationInput:
    """Prepare one case with the exact existing signal biomarker extractor."""
    extractor = biomarker_extractor or existing_biomarker_extractor()
    biomarkers = extractor(
        list(eeg_paths),
        list(emg_paths),
        hand_function_stage=clinical_scores.brunnstrom_hand,
        affected_side=patient.paralysis_side,
        institution=institution,
    )
    return build_evaluation_input(
        patient=patient,
        clinical_scores=clinical_scores,
        biomarkers=biomarkers,
        quality={"status": "pass", "source": "benchmark_single"},
        config=config,
    )


__all__ = [
    "BenchmarkBatch",
    "BenchmarkBatchValidationError",
    "BenchmarkPatientPackage",
    "REQUIRED_SCORE_COLUMNS",
    "prepare_batch_inputs",
    "prepare_single_input",
    "existing_biomarker_extractor",
    "read_benchmark_batch",
    "read_benchmark_batch_zip",
    "read_clinical_scores",
]
